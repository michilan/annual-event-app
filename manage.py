import csv
from pathlib import Path
from typing import Dict, Optional

import click
from openpyxl import load_workbook

from annual_event_app.app import (
    app,
    db,
    Attendee,
    Operator,
    Performance,
    Prize,
    Vote,
    RaffleResult,
    initialize_database,
    parse_bool,
    seed_sample_attendees,
)


@click.group()
def cli():
    """Utility commands for administering the annual event application."""


def _assign_credentials(attendee: Attendee, login_account: Optional[str], password: Optional[str]) -> None:
    desired = (login_account or attendee.login_account or '').strip()
    if not desired:
        desired = str(attendee.id) if attendee.id else ''
    if desired:
        base_account = desired
        unique_account = base_account
        suffix = 1
        while Attendee.query.filter(Attendee.login_account == unique_account, Attendee.id != attendee.id).first():
            unique_account = f"{base_account}{suffix}"
            suffix += 1
        attendee.login_account = unique_account
    if password:
        attendee.set_password(password.strip())
    elif attendee.login_account and not attendee.password_hash:
        attendee.set_password(attendee.login_account)


def _upsert_attendee(data: Dict[str, object]) -> Attendee:
    attendee_id = data.get('id')
    attendee = None
    if attendee_id:
        attendee = Attendee.query.get(attendee_id)
    if attendee is None and data.get('login_account'):
        attendee = Attendee.query.filter_by(login_account=data['login_account']).first()
    if attendee is None:
        attendee = Attendee()
        if attendee_id:
            attendee.id = attendee_id
        db.session.add(attendee)
    attendee.name = data.get('name') or attendee.name or ''
    attendee.call_name = data.get('call_name') or attendee.call_name
    attendee.club_name = data.get('club_name') or attendee.club_name
    attendee.district = data.get('district') or attendee.district
    attendee.event_checked_in = data.get('event_checked_in', attendee.event_checked_in or False)
    attendee.show_checked_in = data.get('show_checked_in', attendee.show_checked_in or False)
    attendee.has_voted = data.get('has_voted', attendee.has_voted or False)
    attendee.has_drawn = data.get('has_drawn', attendee.has_drawn or False)
    if attendee.id is None:
        db.session.flush()
    _assign_credentials(attendee, data.get('login_account'), data.get('password'))
    return attendee


@cli.command()
@click.option('--with-sample', is_flag=True, help='Populate sample attendees after reset.')
def reset_db(with_sample):
    """Clear all records and recreate base data."""
    with app.app_context():
        for model in (Vote, RaffleResult, Prize, Performance, Attendee, Operator):
            model.query.delete()
        db.session.commit()
        initialize_database()
        if with_sample:
            seed_sample_attendees()
        click.echo('Database has been reset.')


@cli.command()
@click.argument('input_path', type=click.Path(exists=True, dir_okay=False))
@click.option('--reset', is_flag=True, help='Clear existing attendees before import.')
def import_attendees(input_path, reset):
    """Import attendees from CSV or Excel (xlsx)."""
    path = Path(input_path)
    with app.app_context():
        if reset:
            Vote.query.delete()
            RaffleResult.query.delete()
            db.session.commit()
            Attendee.query.delete()
            db.session.commit()
        records: list[Dict[str, str]] = []
        suffix = path.suffix.lower()
        if suffix == '.csv':
            with path.open(encoding='utf-8', newline='') as csvfile:
                reader = csv.DictReader(csvfile)
                for row in reader:
                    records.append({(k or '').strip().lower(): (v or '').strip() for k, v in row.items()})
        elif suffix in {'.xlsx', '.xlsm', '.xltx', '.xltm'}:
            workbook = load_workbook(path, data_only=True)
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
            if not rows:
                click.echo('Workbook is empty.')
                return
            headers = [str(h).strip().lower() if h is not None else '' for h in rows[0]]
            for row in rows[1:]:
                if all(cell is None or str(cell).strip() == '' for cell in row):
                    continue
                record: Dict[str, str] = {}
                for header, cell in zip(headers, row):
                    record[header] = '' if cell is None else str(cell).strip()
                records.append(record)
        else:
            click.echo('Unsupported file format. Please provide CSV or XLSX.')
            return

        header_map = {
            'id': 'id',
            '序號': 'id',
            '報到序號': 'id',
            'number': 'id',
            '姓名': 'name',
            'name': 'name',
            'call name': 'call_name',
            '暱稱': 'call_name',
            '所屬社': 'club_name',
            '社別': 'club_name',
            'club': 'club_name',
            '地區': 'district',
            'district': 'district',
            'login_account': 'login_account',
            'account': 'login_account',
            '帳號': 'login_account',
            'password': 'password',
            '密碼': 'password',
            'event_checked_in': 'event_checked_in',
            '第一階段': 'event_checked_in',
            'show_checked_in': 'show_checked_in',
            '第二階段': 'show_checked_in',
            'has_voted': 'has_voted',
            '已評分': 'has_voted',
            'has_drawn': 'has_drawn',
            '已中獎': 'has_drawn',
        }

        def transform(row: Dict[str, str]) -> Dict[str, object]:
            data: Dict[str, object] = {}
            for key, value in row.items():
                normalized = header_map.get(key.lower())
                if not normalized or value == '':
                    continue
                if normalized == 'id':
                    try:
                        data['id'] = int(float(value))
                    except ValueError:
                        pass
                elif normalized in {'event_checked_in', 'show_checked_in', 'has_voted', 'has_drawn'}:
                    data[normalized] = parse_bool(value)
                elif normalized in {'login_account', 'password', 'name', 'call_name', 'club_name', 'district'}:
                    data[normalized] = value
            return data

        count = 0
        for raw in records:
            data = transform(raw)
            if not data.get('name'):
                continue
            _upsert_attendee(data)
            count += 1
        db.session.commit()
        click.echo(f'Imported or updated {count} attendees from {path.name}.')


@cli.command()
@click.option('--default-password', default='club123', show_default=True, help='Password assigned to newly created operators.')
@click.option('--overwrite', is_flag=True, help='Reset password for existing operators to the default password.')
def sync_operators(default_password, overwrite):
    """Create or update club operators based on existing attendees."""
    with app.app_context():
        db.create_all()
        clubs = [row[0] for row in db.session.query(Attendee.club_name).filter(Attendee.club_name.isnot(None)).distinct()]
        created = 0
        updated = 0
        for idx, club in enumerate(clubs, start=1):
            if not club:
                continue
            operator = Operator.query.filter_by(club_name=club).first()
            if operator is None:
                operator = Operator(club_name=club, name=f'{club} 執秘')
                db.session.add(operator)
                created += 1
            else:
                updated += 1
            base = ''.join(ch for ch in club if ch.isascii() and ch.isalnum()) or f'club{idx:03d}'
            base = base.lower()
            unique = base
            suffix = 1
            with db.session.no_autoflush:
                while Operator.query.filter(Operator.login_account == unique, Operator.id != operator.id).first():
                    unique = f"{base}{suffix}"
                    suffix += 1
            operator.login_account = unique
            if overwrite or not operator.password_hash:
                operator.set_password(default_password)
        db.session.commit()
        click.echo(f'Synced operators. Created: {created}, existing updated: {updated}.')

if __name__ == '__main__':
    cli()
